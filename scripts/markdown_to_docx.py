#!/usr/bin/env python3
"""
Convert a simple Markdown file into a .docx document.

Supported Markdown features:
- #, ##, ###, #### headings
- plain paragraphs
- hard line breaks inside paragraphs
- flat bullet lists
- flat numbered lists
- blockquotes
- simple inline bold / italic
- simple pipe tables

This utility is intentionally narrow and predictable for grant attachments.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor
from docx.text.paragraph import Paragraph


HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_RE = re.compile(r"^-\s+(.*)$")
NUMBER_RE = re.compile(r"^\d+\.\s+(.*)$")
TABLE_RE = re.compile(r"^\|.*\|$")
BLOCKQUOTE_RE = re.compile(r"^>\s?(.*)$")
INLINE_TOKEN_RE = re.compile(r"(`[^`\n]+`|\*\*[^*\n]+\*\*|\*[^*\n]+\*)")
ACCENT_BLUE = RGBColor.from_string("1E5AA8")


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("1E1E2E")


def apply_run_style(run, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("1E1E2E")
    run.bold = bold
    run.italic = italic


def append_markdown_runs(
    paragraph: Paragraph,
    text: str,
    *,
    color: RGBColor | None = None,
    font_name: str = "Arial",
    size_pt: float = 11,
) -> None:
    pieces = INLINE_TOKEN_RE.split(text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`") and len(piece) >= 3:
            run = paragraph.add_run(piece[1:-1])
            apply_run_style(run)
            run.font.name = font_name
            run.font.size = Pt(size_pt)
            if color:
                run.font.color.rgb = color
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) >= 4:
            content = piece[2:-2]
            if content.startswith("`") and content.endswith("`") and len(content) >= 3:
                content = content[1:-1]
            run = paragraph.add_run(content)
            apply_run_style(run, bold=True)
            run.font.name = font_name
            run.font.size = Pt(size_pt)
            if color:
                run.font.color.rgb = color
            continue
        if piece.startswith("*") and piece.endswith("*") and len(piece) >= 3:
            content = piece[1:-1]
            if content.startswith("`") and content.endswith("`") and len(content) >= 3:
                content = content[1:-1]
            run = paragraph.add_run(content)
            apply_run_style(run, italic=True)
            run.font.name = font_name
            run.font.size = Pt(size_pt)
            if color:
                run.font.color.rgb = color
            continue
        run = paragraph.add_run(piece)
        apply_run_style(run)
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        if color:
            run.font.color.rgb = color


def add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    for idx, line in enumerate(text.split("\n")):
        if idx > 0:
            p.add_run().add_break()
        append_markdown_runs(p, line)


def add_heading(doc: Document, level: int, text: str) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    size = max(12, 18 - (level - 1) * 2)
    append_markdown_runs(p, text, color=ACCENT_BLUE, font_name="Georgia", size_pt=size)
    for run in p.runs:
        run.bold = True
        run.font.name = "Georgia"
        run.font.size = Pt(size)
        run.font.color.rgb = ACCENT_BLUE
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)


def add_blockquote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    for idx, line in enumerate(text.split("\n")):
        if idx > 0:
            p.add_run().add_break()
        append_markdown_runs(p, line)
    for run in p.runs:
        run.italic = True


def split_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip().replace("<br>", "\n") for cell in line.split("|")]


def is_separator_row(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_table(doc: Document, lines: list[str]) -> None:
    rows = [split_table_row(line) for line in lines]
    if len(rows) < 2 or not is_separator_row(lines[1]):
        for row in rows:
            add_paragraph(doc, " | ".join(row))
        return

    header = rows[0]
    data_rows = rows[2:]
    cols = len(header)

    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    for idx, value in enumerate(header):
        cell = table.rows[0].cells[idx]
        paragraph = cell.paragraphs[0]
        append_markdown_runs(paragraph, value)

    for row in data_rows:
        tr = table.add_row().cells
        for idx in range(cols):
            paragraph = tr[idx].paragraphs[0]
            append_markdown_runs(paragraph, row[idx] if idx < len(row) else "")


def join_paragraph_lines(lines: list[str]) -> str:
    if not lines:
        return ""

    parts: list[str] = []
    for idx, line in enumerate(lines):
        stripped = line.strip()
        if idx == 0:
            parts.append(stripped)
            continue

        prev = lines[idx - 1]
        if prev.endswith("  "):
            parts.append("\n")
        else:
            parts.append(" ")
        parts.append(stripped)
    return "".join(parts)


def add_image_with_caption(
    doc: Document,
    image_path: Path,
    caption: str,
    *,
    width_cm: float = 12.5,
) -> None:
    if not image_path.exists():
        return

    pic_paragraph = doc.add_paragraph()
    pic_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic_paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))

    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_run = caption_paragraph.add_run(caption)
    caption_run.font.name = "Arial"
    caption_run.font.size = Pt(10.5)
    caption_run.font.bold = True
    caption_run.font.color.rgb = ACCENT_BLUE


def load_character_image_blocks(manifest_path: Path) -> list[tuple[str, list[str]]]:
    raw = json.loads(manifest_path.read_text(encoding="utf-8"))
    if not isinstance(raw, list):
        raise ValueError("Character image manifest must be a JSON array")

    blocks: list[tuple[str, list[str]]] = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        caption = str(item.get("caption", "")).strip()
        images = item.get("images")
        if not caption or not isinstance(images, list):
            continue
        normalized = [str(image).strip() for image in images if str(image).strip()]
        if normalized:
            blocks.append((caption, normalized))
    return blocks


def append_character_reference_images(
    doc: Document,
    blocks: list[tuple[str, list[str]]],
) -> None:
    add_heading(doc, 2, "Визуальные ориентиры персонажей")
    intro = (
        "Ниже приведены визуальные ориентиры по ключевым персонажам и отношениям "
        "внутри истории. Они используются как референс по типажам, пластике, "
        "возрастному ощущению и эмоциональному тону сцен."
    )
    add_paragraph(doc, intro)

    for caption, image_paths in blocks:
        subheading = doc.add_paragraph()
        subheading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subheading.paragraph_format.space_before = Pt(6)
        subheading.paragraph_format.space_after = Pt(4)
        run = subheading.add_run(caption)
        run.font.name = "Georgia"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = ACCENT_BLUE
        for image_path in image_paths:
            add_image_with_caption(doc, Path(image_path), caption)


def parse_blocks(text: str) -> list[tuple[str, object]]:
    lines = text.splitlines()
    blocks: list[tuple[str, object]] = []
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            blocks.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
            i += 1
            continue

        blockquote = BLOCKQUOTE_RE.match(stripped)
        if blockquote:
            quote_lines = []
            while i < len(lines):
                current = lines[i].rstrip("\n")
                match = BLOCKQUOTE_RE.match(current.strip())
                if not match:
                    break
                quote_lines.append(match.group(1))
                i += 1
            blocks.append(("blockquote", join_paragraph_lines(quote_lines)))
            continue

        if TABLE_RE.match(stripped):
            table_lines = []
            while i < len(lines) and TABLE_RE.match(lines[i].strip()):
                table_lines.append(lines[i].strip())
                i += 1
            blocks.append(("table", table_lines))
            continue

        bullet = BULLET_RE.match(stripped)
        if bullet:
            items = []
            while i < len(lines):
                current = lines[i].strip()
                match = BULLET_RE.match(current)
                if not match:
                    break
                items.append(match.group(1).strip())
                i += 1
            blocks.append(("bullet_list", items))
            continue

        number = NUMBER_RE.match(stripped)
        if number:
            items = []
            while i < len(lines):
                current = lines[i].strip()
                match = NUMBER_RE.match(current)
                if not match:
                    break
                items.append(match.group(1).strip())
                i += 1
            blocks.append(("number_list", items))
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            current_raw = lines[i]
            current = current_raw.strip()
            if not current:
                i += 1
                break
            if (
                HEADING_RE.match(current)
                or TABLE_RE.match(current)
                or BULLET_RE.match(current)
                or NUMBER_RE.match(current)
                or BLOCKQUOTE_RE.match(current)
            ):
                break
            paragraph_lines.append(current_raw)
            i += 1
        blocks.append(("paragraph", join_paragraph_lines(paragraph_lines)))

    return blocks


def markdown_to_docx(
    input_path: Path,
    output_path: Path,
    *,
    character_image_manifest: Path | None = None,
) -> None:
    text = input_path.read_text(encoding="utf-8")
    blocks = parse_blocks(text)

    doc = Document()
    configure_document(doc)

    for block_type, value in blocks:
        if block_type == "heading":
            level, text_value = value
            add_heading(doc, level, text_value)
        elif block_type == "paragraph":
            add_paragraph(doc, value)
        elif block_type == "bullet_list":
            for item in value:
                add_paragraph(doc, item, style="List Bullet")
        elif block_type == "number_list":
            for item in value:
                add_paragraph(doc, item, style="List Number")
        elif block_type == "blockquote":
            add_blockquote(doc, value)
        elif block_type == "table":
            add_table(doc, value)

    if character_image_manifest is not None:
        append_character_reference_images(
            doc,
            load_character_image_blocks(character_image_manifest),
        )

    doc.save(str(output_path))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert a simple Markdown file into .docx."
    )
    parser.add_argument("input", help="Path to the source .md file")
    parser.add_argument(
        "-o",
        "--output",
        help="Optional output .docx path. Defaults next to the source file.",
    )
    parser.add_argument(
        "--character-image-manifest",
        help="Optional JSON manifest with image blocks to append.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = Path(args.input).expanduser().resolve()

    if not input_path.exists():
        print(f"Input file does not exist: {input_path}", file=sys.stderr)
        return 1

    if input_path.suffix.lower() != ".md":
        print(f"Expected a .md file, got: {input_path.name}", file=sys.stderr)
        return 1

    output_path = (
        Path(args.output).expanduser().resolve()
        if args.output
        else input_path.with_suffix(".docx")
    )

    character_image_manifest = (
        Path(args.character_image_manifest).expanduser().resolve()
        if args.character_image_manifest
        else None
    )
    markdown_to_docx(
        input_path,
        output_path,
        character_image_manifest=character_image_manifest,
    )
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
